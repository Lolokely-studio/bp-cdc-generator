"""Produit `app/export/templates/model.docx`, le modèle Word par défaut.

Un modèle fait à la main dans Word peut le remplacer : il lui suffit de porter
les trois balises ci-dessous et les styles `Heading 1`, `Heading 2`,
`List Bullet`, `Caption` et `Table Grid`, que le rendu emploie.

    uv run python scripts/build_word_template.py
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

OUTPUT = Path(__file__).resolve().parents[1] / "app/export/templates/model.docx"


def build() -> None:
    document = Document()
    base = document.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("{{ title }}")
    run.bold = True
    run.font.size = Pt(24)

    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice_run = notice.add_run("{{ draft_notice }}")
    notice_run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)

    document.add_page_break()
    document.add_paragraph("{{p body }}")

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = "{{ title }}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(f"modèle écrit : {OUTPUT}")
