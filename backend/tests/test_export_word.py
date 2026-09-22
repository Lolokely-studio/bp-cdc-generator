import io

from docx import Document

from app.agent.state import BulletList, Paragraph, Placeholder, Table
from app.export.charts import Chart
from app.export.document import ExportDocument, ExportSection
from app.export.word import render_word


def _png():
    import matplotlib.pyplot as plt
    from app.export import charts  # impose le moteur Agg

    figure, axis = plt.subplots()
    axis.plot([1, 2])
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png")
    plt.close(figure)
    return buffer.getvalue()


def _document(draft=False, missing=()):
    return ExportDocument(
        document="bp", title="Business plan — CoachDom", draft=draft,
        missing=list(missing),
        sections=[
            ExportSection("Résumé", [Paragraph(text="Un résumé clair.")]),
            ExportSection("Finances", [
                Table(number=1, title="Compte de résultat", columns=["Année", "CA"],
                      rows=[["1", "100 000,00 €"]]),
                BulletList(items=["premier point", "second point"]),
                Placeholder(label="taux d'emprunt"),
            ]),
        ],
    )


def _read(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def test_the_word_file_carries_every_section_and_block():
    word = _read(render_word(_document(), []))
    text = "\n".join(p.text for p in word.paragraphs)
    assert "Business plan — CoachDom" in text
    assert "Résumé" in text and "Finances" in text
    assert "Un résumé clair." in text
    assert "premier point" in text
    assert len(word.tables) == 1
    assert word.tables[0].cell(1, 1).text == "100 000,00 €"


def test_tables_carry_their_document_number_in_the_caption():
    word = _read(render_word(_document(), []))
    captions = [p.text for p in word.paragraphs if p.text.startswith("Tableau ")]
    assert captions == ["Tableau 1 — Compte de résultat"]


def test_a_draft_says_so_and_a_final_document_does_not():
    draft = "\n".join(p.text for p in _read(render_word(_document(draft=True), [])).paragraphs)
    final = "\n".join(p.text for p in _read(render_word(_document(draft=False), [])).paragraphs)
    assert "Brouillon" in draft
    assert "Brouillon" not in final


def test_missing_data_are_gathered_in_an_appendix():
    word = _read(render_word(_document(missing=["taux d'emprunt"]), []))
    text = [p.text for p in word.paragraphs]
    assert "Données à compléter" in text
    assert "taux d'emprunt" in "\n".join(text)


def test_no_appendix_when_nothing_is_missing():
    word = _read(render_word(_document(missing=[]), []))
    assert "Données à compléter" not in [p.text for p in word.paragraphs]


def test_charts_are_embedded_as_images():
    word = _read(render_word(_document(), [Chart("Trésorerie", _png())]))
    assert len(word.inline_shapes) == 1


def test_no_template_tag_survives_the_rendering():
    """Une balise restée telle quelle est un modèle et un rendu qui ne se
    sont pas compris : on la verrait imprimée dans le document."""
    word = _read(render_word(_document(), []))
    text = "\n".join(p.text for p in word.paragraphs)
    footer = word.sections[0].footer.paragraphs[0].text
    assert "{{" not in text and "{{" not in footer
