import io
from pathlib import Path

from docx.shared import Cm
from docxtpl import DocxTemplate
from jinja2 import Environment, StrictUndefined

from app.agent.state import BulletList, Paragraph, Placeholder, Table
from app.export.charts import Chart
from app.export.document import ExportDocument

MODEL = Path(__file__).parent / "templates" / "model.docx"

DRAFT_NOTICE = ("Brouillon — certaines sections ne sont pas encore validées. "
                "Ce document n'est pas définitif.")

# `autoescape` : le titre porte le nom du projet, du texte libre, substitué
# tel quel dans le XML du document ; un « & » le rendait invalide.
# `StrictUndefined` : une balise mal orthographiée dans un modèle refait à la
# main lève à l'export au lieu de rendre un vide que personne ne remarquerait.
_JINJA = Environment(undefined=StrictUndefined, autoescape=True)


def _add_table(subdoc, block: Table) -> None:
    subdoc.add_paragraph(f"Tableau {block.number} — {block.title}",
                         style="Caption")
    table = subdoc.add_table(rows=1 + len(block.rows), cols=len(block.columns))
    table.style = "Table Grid"
    for index, heading in enumerate(block.columns):
        cell = table.cell(0, index)
        cell.text = heading
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row_index, row in enumerate(block.rows, start=1):
        for index, value in enumerate(row):
            table.cell(row_index, index).text = value


def render_word(document: ExportDocument, charts: list[Chart]) -> bytes:
    """Le Word, à partir du modèle et d'un sous-document dans ses styles.

    Le modèle ne connaît que trois balises : `title`, `draft_notice` et
    `body`. Tout le contenu passe par le sous-document, ce qui laisse le
    modèle remplaçable par un fichier refait à la main dans Word.
    """
    template = DocxTemplate(MODEL)
    body = template.new_subdoc()

    for section in document.sections:
        body.add_heading(section.title, level=1)
        for block in section.blocks:
            if isinstance(block, Paragraph):
                body.add_paragraph(block.text)
            elif isinstance(block, Table):
                _add_table(body, block)
            elif isinstance(block, BulletList):
                for item in block.items:
                    body.add_paragraph(item, style="List Bullet")
            elif isinstance(block, Placeholder):
                # Visible dans le texte, rassemblé en annexe : le lecteur
                # voit le trou là où il est, et la liste complète à la fin.
                body.add_paragraph(f"[Donnée à compléter : {block.label}]")

    if charts:
        body.add_heading("Graphiques", level=1)
        for chart in charts:
            body.add_paragraph(chart.title, style="Caption")
            body.add_picture(io.BytesIO(chart.png), width=Cm(15))

    if document.missing:
        body.add_heading("Données à compléter", level=1)
        for label in document.missing:
            body.add_paragraph(label, style="List Bullet")

    template.render({
        "title": document.title,
        "draft_notice": DRAFT_NOTICE if document.draft else "",
        "body": body,
    }, jinja_env=_JINJA, autoescape=True)
    output = io.BytesIO()
    template.save(output)
    return output.getvalue()
